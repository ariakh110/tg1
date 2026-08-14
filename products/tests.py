from django.contrib.auth import get_user_model
from io import BytesIO
import datetime
import shutil
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from PIL import Image
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import RoleCode, UserRole

from .models import (
    DeliveryLocation,
    Offer,
    PricingTier,
    Product,
    ProductAuditLog,
    ProductAttributeOption,
    ProductCategory,
    ProductImage,
    ProductSpecification,
    Seller,
)
from .serializers import ProductPageContentWriteSerializer, ProductSpecificationSerializer

User = get_user_model()


def make_test_category_icon(filename="category.png", image_format="PNG"):
    content = BytesIO()
    Image.new("RGB", (24, 24), color=(25, 95, 145)).save(content, format=image_format)
    return SimpleUploadedFile(filename, content.getvalue(), content_type=f"image/{image_format.lower()}")


class SellerProfileTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="seller_user", password="pass1234")

    def test_create_seller_profile_is_idempotent(self):
        self.client.force_authenticate(self.user)
        payload = {
            "company_name": "Alpha Steel",
            "business_type": "Trading",
            "location": "Tehran",
        }
        first = self.client.post("/api/sellers/", payload, format="json")
        self.assertEqual(first.status_code, status.HTTP_201_CREATED)
        seller_id = first.data["id"]

        second = self.client.post("/api/sellers/", payload, format="json")
        self.assertEqual(second.status_code, status.HTTP_200_OK)
        self.assertEqual(second.data["id"], seller_id)


class ProductAssetPermissionTests(APITestCase):
    def setUp(self):
        self.seller_owner_user = User.objects.create_user(
            username="seller_owner", password="pass1234"
        )
        self.seller_other_user = User.objects.create_user(
            username="seller_other", password="pass1234"
        )

        UserRole.objects.update_or_create(
            user=self.seller_owner_user,
            role=RoleCode.SELLER,
            defaults={"is_active": True, "activated_at": timezone.now()},
        )
        UserRole.objects.update_or_create(
            user=self.seller_other_user,
            role=RoleCode.SELLER,
            defaults={"is_active": True, "activated_at": timezone.now()},
        )

        self.owner_seller = Seller.objects.create(
            user=self.seller_owner_user,
            company_name="Owner Co",
            business_type="Trading",
            location="Isfahan",
            is_verified=True,
        )
        self.other_seller = Seller.objects.create(
            user=self.seller_other_user,
            company_name="Other Co",
            business_type="Trading",
            location="Shiraz",
            is_verified=True,
        )

        category = ProductCategory.objects.create(name="Permission Steel")
        self.product = Product.objects.create(
            category=category,
            name="Hot Rolled Coil",
            short_description="HRC",
            description="HRC product",
        )
        Offer.objects.create(product=self.product, seller=self.owner_seller, is_active=True)

    def test_only_offer_owner_can_upload_product_image(self):
        with tempfile.TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            self.client.force_authenticate(self.seller_other_user)
            forbidden = self.client.post(
                "/api/product-images/",
                {"product": self.product.id, "image": make_test_category_icon("forbidden.png")},
                format="multipart",
            )
            self.assertEqual(forbidden.status_code, status.HTTP_403_FORBIDDEN)

            self.client.force_authenticate(self.seller_owner_user)
            allowed = self.client.post(
                "/api/product-images/",
                {"product": self.product.id, "image": make_test_category_icon("allowed.png")},
                format="multipart",
            )
            self.assertEqual(allowed.status_code, status.HTTP_201_CREATED)
            self.assertTrue(allowed.data["is_featured"])

    def test_featured_product_image_lifecycle_is_deterministic(self):
        with tempfile.TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            self.client.force_authenticate(self.seller_owner_user)
            first = self.client.post(
                "/api/product-images/",
                {"product": self.product.id, "image": make_test_category_icon("first.png")},
                format="multipart",
            )
            second = self.client.post(
                "/api/product-images/",
                {"product": self.product.id, "image": make_test_category_icon("second.png")},
                format="multipart",
            )

            self.assertEqual(first.status_code, status.HTTP_201_CREATED, first.data)
            self.assertEqual(second.status_code, status.HTTP_201_CREATED, second.data)
            self.assertTrue(first.data["is_featured"])
            self.assertFalse(second.data["is_featured"])

            selected = self.client.patch(
                f"/api/product-images/{second.data['id']}/",
                {"is_featured": True},
                format="json",
            )
            self.assertEqual(selected.status_code, status.HTTP_200_OK, selected.data)
            self.assertFalse(ProductImage.objects.get(pk=first.data["id"]).is_featured)
            self.assertTrue(ProductImage.objects.get(pk=second.data["id"]).is_featured)

            listing = self.client.get("/api/product-images/", {"product": self.product.id})
            self.assertEqual(listing.status_code, status.HTTP_200_OK, listing.data)
            listed_images = listing.data.get("results", listing.data)
            self.assertEqual(listed_images[0]["id"], second.data["id"])

            deleted = self.client.delete(f"/api/product-images/{second.data['id']}/")
            self.assertEqual(deleted.status_code, status.HTTP_204_NO_CONTENT)
            self.assertTrue(ProductImage.objects.get(pk=first.data["id"]).is_featured)


class ProductTaxonomyTests(APITestCase):
    def make_product(self, category, name="ورق تست"):
        return Product.objects.create(
            category=category,
            name=name,
            short_description=name,
            description=name,
        )

    def test_category_endpoint_exposes_backend_taxonomy_metadata(self):
        black_sheet = ProductCategory.objects.get(code="sheet-black")

        res = self.client.get("/api/categories/", {"code": "sheet-black"})

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        item = res.data["results"][0]
        self.assertEqual(item["name"], "ورق سیاه")
        self.assertEqual(item["resolved_product_kind"], "sheet")
        self.assertEqual(item["merged_spec_defaults"]["steel_grade"], "ST37")
        self.assertIn("thickness_mm", item["merged_required_spec_fields"])
        self.assertEqual(item["children_count"], black_sheet.get_children().count())

    def test_spec_serializer_applies_category_defaults_and_required_fields(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = self.make_product(category, "ورق ۲ میل مبارکه")

        missing = ProductSpecificationSerializer(
            data={"product": product.id, "width_mm": "1000"}
        )
        self.assertFalse(missing.is_valid())
        self.assertIn("thickness_mm", missing.errors)

        serializer = ProductSpecificationSerializer(
            data={
                "product": product.id,
                "thickness_mm": "2",
                "width_mm": "1000",
                "length_mm": "6000",
                "factory": "mobarakeh",
                "cut_type": "cut",
            }
        )
        self.assertTrue(serializer.is_valid(), serializer.errors)
        spec = serializer.save()

        self.assertEqual(spec.material_type, "sheet")
        self.assertEqual(spec.steel_grade, "ST37")
        self.assertEqual(spec.surface_finish, "black")
        self.assertEqual(spec.manufacturing_process, "sheet")
        self.assertEqual(spec.factory, "mobarakeh")
        self.assertEqual(spec.cut_type, "cut")

        patch = ProductSpecificationSerializer(
            spec,
            data={"width_mm": "1250"},
            partial=True,
        )
        self.assertTrue(patch.is_valid(), patch.errors)

    def test_non_sheet_category_can_define_specs_without_steel_grade(self):
        category = ProductCategory.objects.get(code="profile-box")
        product = self.make_product(category, "قوطی تست")

        serializer = ProductSpecificationSerializer(
            data={
                "product": product.id,
                "width_mm": "40",
                "height_mm": "40",
                "thickness_mm": "2",
            }
        )

        self.assertTrue(serializer.is_valid(), serializer.errors)
        spec = serializer.save()
        self.assertEqual(spec.material_type, "profile")
        self.assertEqual(spec.steel_grade, "")

    def test_product_filters_include_category_descendants_and_spec_defaults(self):
        mobarakeh = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        sheet = ProductCategory.objects.get(code="sheet")
        oiled = ProductCategory.objects.get(code="sheet-oiled")
        black_product = self.make_product(mobarakeh, "ورق سیاه مبارکه تست")
        parent_registered_product = self.make_product(sheet, "ورق سیاه ثبت شده روی دسته مادر")
        oiled_product = self.make_product(oiled, "ورق روغنی تست")
        ProductSpecification.objects.create(
            product=black_product,
            material_type="sheet",
            steel_grade="ST37",
            surface_finish="black",
            manufacturing_process="sheet",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm="2",
            width_mm="1000",
            length_mm="6000",
        )
        ProductSpecification.objects.create(
            product=parent_registered_product,
            material_type="sheet",
            steel_grade="ST37",
            surface_finish="black",
            manufacturing_process="sheet",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm="2",
            width_mm="1250",
            length_mm="6000",
        )
        ProductSpecification.objects.create(
            product=oiled_product,
            material_type="sheet",
            steel_grade="ST12",
            surface_finish="oiled",
            manufacturing_process="sheet",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm="1",
            width_mm="1000",
            length_mm="6000",
        )

        res = self.client.get(
            "/api/products/",
            {"category_code": "sheet-black", "steel_grade": "ST37", "surface_finish": "black"},
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        names = {item["name"] for item in res.data["results"]}
        self.assertIn(black_product.name, names)
        self.assertIn(parent_registered_product.name, names)
        self.assertNotIn(oiled_product.name, names)

    def test_public_product_list_excludes_inactive_products(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        active = self.make_product(category, "ورق فعال")
        inactive = Product.objects.create(
            category=category,
            name="ورق غیر فعال",
            short_description="ورق غیر فعال",
            description="ورق غیر فعال",
            is_active=False,
        )

        res = self.client.get("/api/products/", {"category_code": "sheet-black"})

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        ids = {item["id"] for item in res.data["results"]}
        self.assertIn(active.id, ids)
        self.assertNotIn(inactive.id, ids)

    def test_active_category_endpoint_exposes_only_matching_product_types(self):
        root = ProductCategory.objects.create(
            name="دسته تست",
            code="test-root",
            product_kind="sheet",
            spec_defaults={"material_type": "sheet"},
        )
        active_category = ProductCategory.objects.create(
            name="دسته فعال تست",
            code="test-active",
            parent=root,
            product_kind="sheet",
            spec_defaults={"surface_finish": "black"},
        )
        inactive_category = ProductCategory.objects.create(
            name="دسته غیر فعال تست",
            code="test-inactive",
            parent=root,
            product_kind="sheet",
            spec_defaults={"surface_finish": "oiled"},
        )
        unrelated_root = ProductCategory.objects.create(name="دسته بی محصول", code="test-empty-root")
        product_on_root = self.make_product(root, "ورق فعال دسته مادر")
        ProductSpecification.objects.create(
            product=product_on_root,
            material_type="sheet",
            steel_grade="ST37",
            surface_finish="black",
            manufacturing_process="sheet",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm="2",
            width_mm="1000",
            length_mm="6000",
        )
        Product.objects.create(
            category=inactive_category,
            name="ورق غیرفعال دسته",
            short_description="ورق غیرفعال دسته",
            description="ورق غیرفعال دسته",
            is_active=False,
        )

        res = self.client.get("/api/categories/active-with-products/")

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        codes = {item["code"] for item in res.data["results"]}
        self.assertIn("test-root", codes)
        self.assertIn("test-active", codes)
        self.assertNotIn("test-inactive", codes)
        self.assertNotIn(unrelated_root.code, codes)

    def test_active_category_endpoint_includes_acid_washed_sheet_when_matching_product_exists(self):
        sheet = ProductCategory.objects.get(code="sheet")
        product = self.make_product(sheet, "ورق اسیدشویی تست")
        ProductSpecification.objects.create(
            product=product,
            material_type="sheet",
            steel_grade="ST37",
            surface_finish="acid_washed",
            manufacturing_process="sheet",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm="2",
            width_mm="1250",
            length_mm="6000",
        )

        res = self.client.get("/api/categories/active-with-products/")

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        codes = {item["code"] for item in res.data["results"]}
        self.assertIn("sheet", codes)
        self.assertIn("sheet-acid-washed", codes)

    def test_sheet_grade_options_depend_on_sheet_type(self):
        black = ProductAttributeOption.objects.get(group="surface_finish", value="black")
        alloy = ProductAttributeOption.objects.get(group="surface_finish", value="alloy")

        black_res = self.client.get("/api/attribute-options/", {"group": "steel_grade", "parent": black.id})
        alloy_res = self.client.get("/api/attribute-options/", {"group": "steel_grade", "parent": alloy.id})

        self.assertEqual(black_res.status_code, status.HTTP_200_OK)
        self.assertEqual(alloy_res.status_code, status.HTTP_200_OK)
        black_values = {item["value"] for item in black_res.data["results"]}
        alloy_values = {item["value"] for item in alloy_res.data["results"]}
        self.assertIn("ST37", black_values)
        self.assertNotIn("CK45", black_values)
        self.assertNotIn("ST52", black_values)
        self.assertIn("CK45", alloy_values)
        self.assertIn("ST52", alloy_values)

    def test_sheet_spec_rejects_grade_not_allowed_for_selected_sheet_type(self):
        sheet = ProductCategory.objects.get(code="sheet")
        product = self.make_product(sheet, "ورق سیاه با گرید نامعتبر")

        serializer = ProductSpecificationSerializer(
            data={
                "product": product.id,
                "material_type": "sheet",
                "steel_grade": "CK45",
                "surface_finish": "black",
                "manufacturing_process": "sheet",
                "factory": "mobarakeh",
                "cut_type": "cut",
                "thickness_mm": "2",
                "width_mm": "1000",
                "length_mm": "6000",
            }
        )

        self.assertFalse(serializer.is_valid())
        self.assertIn("steel_grade", serializer.errors)

    def test_billet_category_and_grades_are_seeded(self):
        billet = ProductCategory.objects.get(code="billet")
        grades = set(
            ProductAttributeOption.objects.filter(group="steel_grade", product_kind="billet")
            .values_list("value", flat=True)
        )

        self.assertEqual(billet.resolved_product_kind(), "billet")
        self.assertIn("3SP", grades)
        self.assertIn("5SP", grades)

    def test_product_list_returns_min_price_and_origin_for_market_rows(self):
        user = User.objects.create_user(username="origin_seller", password="pass1234")
        seller = Seller.objects.create(
            user=user,
            company_name="فولاد مبارکه",
            business_type="Factory",
            location="اصفهان",
            is_verified=True,
        )
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = self.make_product(category, "ورق سیاه ۲ میل مبارکه")
        ProductSpecification.objects.create(
            product=product,
            material_type="sheet",
            steel_grade="ST37",
            surface_finish="black",
            manufacturing_process="sheet",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm="2",
            width_mm="1000",
            length_mm="6000",
        )
        offer = Offer.objects.create(product=product, seller=seller)
        PricingTier.objects.create(
            offer=offer,
            tier_name="قیمت روز",
            unit_price="42000",
            minimum_quantity=1,
        )
        DeliveryLocation.objects.create(
            offer=offer,
            incoterm="EXW",
            country="ایران",
            province="اصفهان",
            city="مبارکه",
            address="کارخانه",
        )

        res = self.client.get("/api/products/", {"category": category.parent_id})

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        item = next(row for row in res.data["results"] if row["id"] == product.id)
        self.assertEqual(item["min_price"], "42000.00")
        delivery = item["offers"][0]["delivery_options"][0]
        self.assertEqual(delivery["province"], "اصفهان")
        self.assertEqual(delivery["city"], "مبارکه")
        self.assertEqual(delivery["address"], "کارخانه")

    def test_attribute_options_endpoint_filters_backend_combo_options(self):
        province = ProductAttributeOption.objects.get(group="province", value="isfahan")

        res = self.client.get(
            "/api/attribute-options/",
            {"group": "city", "parent": province.id},
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        labels = {item["label"] for item in res.data["results"]}
        self.assertIn("مبارکه", labels)


class AdminProductImportTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="product_admin",
            password="pass1234",
            is_staff=True,
        )
        self.seller_user = User.objects.create_user(username="price_seller", password="pass1234")
        self.seller = Seller.objects.create(
            user=self.seller_user,
            company_name="کاوکس",
            business_type="Trading",
            location="تهران",
            is_verified=True,
        )

    def test_admin_can_create_product_with_price_and_origin(self):
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "ورق سیاه ۲ میل مبارکه ادمین",
                "category_code": "sheet",
                "seller_id": self.seller.id,
                "price": "45000",
                "price_basis": "kg",
                "condition_label": "عرض 1000 طول 6000",
                "dimension_width_mm": "1000",
                "dimension_length_mm": "6000",
                "steel_grade": "ST37",
                "manufacturing_process": "sheet",
                "surface_finish": "black",
                "factory": "mobarakeh",
                "cut_type": "cut",
                "thickness_mm": "2",
                "width_mm": "1000",
                "length_mm": "6000",
                "province": "اصفهان",
                "city": "مبارکه",
                "address": "کارخانه",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        product = Product.objects.get(id=res.data["product_id"])
        self.assertEqual(product.category.code, "sheet")
        self.assertEqual(product.specifications.steel_grade, "ST37")
        self.assertEqual(product.specifications.factory, "mobarakeh")
        offer = product.offers.get(seller=self.seller)
        tier = offer.pricing_tiers.get()
        self.assertIsNotNone(tier.price_verified_at)
        self.assertEqual(tier.price_basis, "kg")
        self.assertEqual(tier.condition_label, "عرض 1000 طول 6000")
        self.assertEqual(str(tier.dimension_width_mm), "1000.00")
        self.assertEqual(str(tier.dimension_length_mm), "6000.00")
        self.assertEqual(str(offer.pricing_tiers.get(tier_name="قیمت روز").unit_price), "45000.00")
        self.assertEqual(offer.delivery_options.get().city, "مبارکه")

    def test_admin_can_save_product_page_content_and_related_products(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = Product.objects.create(category=category, name="SEO product")
        related = Product.objects.create(category=category, name="Related product")
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "product_id": product.id,
                "name": product.name,
                "page_content": {
                    "seo_title": "  Dedicated SEO title  ",
                    "meta_description": "Dedicated meta description",
                    "page_h1": "Dedicated H1",
                    "short_description": "Dedicated short description",
                    "description": (
                        '<h1 onclick="alert(1)">Unsafe heading</h1>'
                        '<script>alert(1)</script>'
                        '<p><a href="javascript:alert(1)">Safe text</a></p>'
                    ),
                    "seo_faqs": [{"question": "Question?", "answer": "Answer."}],
                    "seo_index_mode": "index",
                    "related_product_ids": [related.id],
                },
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        product.refresh_from_db()
        self.assertEqual(product.seo_title, "Dedicated SEO title")
        self.assertEqual(product.page_h1, "Dedicated H1")
        self.assertEqual(product.seo_index_mode, Product.SEO_INDEX_INDEX)
        self.assertEqual(product.seo_faqs, [{"question": "Question?", "answer": "Answer."}])
        self.assertEqual(list(product.related_products.values_list("id", flat=True)), [related.id])
        self.assertIn("<h2>Unsafe heading</h2>", product.description)
        self.assertNotIn("script", product.description.lower())
        self.assertNotIn("javascript", product.description.lower())
        self.assertNotIn("onclick", product.description.lower())

        image_serializer = ProductPageContentWriteSerializer(
            product,
            data={
                "description": (
                    '<figure><img src="/media/products/sheet.webp" alt="Sheet" '
                    'onerror="alert(1)" loading="lazy"><figcaption>Product image</figcaption></figure>'
                ),
            },
            partial=True,
        )
        self.assertTrue(image_serializer.is_valid(), image_serializer.errors)
        image_serializer.save()
        self.assertIn('src="/media/products/sheet.webp"', product.description)
        self.assertIn("<figcaption>Product image</figcaption>", product.description)
        self.assertNotIn("onerror", product.description)

        Product.objects.filter(pk=product.pk).update(
            description='<h1>Legacy H1</h1><script>alert(1)</script><p>Legacy safe content</p>',
        )
        legacy_detail = self.client.get(f"/api/products/{product.slug}/")
        self.assertEqual(legacy_detail.status_code, status.HTTP_200_OK, legacy_detail.data)
        self.assertIn("<h2>Legacy H1</h2>", legacy_detail.data["description"])
        self.assertNotIn("script", legacy_detail.data["description"].lower())

        detail = self.client.get(f"/api/products/{product.slug}/")
        self.assertEqual(detail.status_code, status.HTTP_200_OK, detail.data)
        self.assertEqual(detail.data["seo_title"], "Dedicated SEO title")
        self.assertEqual(detail.data["related_products"][0]["id"], related.id)
        self.assertIn("suggested_products", detail.data)

    def test_page_content_serializer_rejects_self_relation_and_incomplete_faq(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = Product.objects.create(category=category, name="Validation product")

        serializer = ProductPageContentWriteSerializer(
            product,
            data={
                "related_product_ids": [product.id],
                "seo_faqs": [{"question": "Only question", "answer": ""}],
            },
            partial=True,
        )

        self.assertFalse(serializer.is_valid())
        self.assertIn("related_product_ids", serializer.errors)
        self.assertIn("seo_faqs", serializer.errors)

    def test_price_only_update_preserves_all_page_content_fields(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        related = Product.objects.create(category=category, name="Preserved related")
        product = Product.objects.create(
            category=category,
            name="Preserved content product",
            short_description="Keep short",
            description="<p>Keep full content</p>",
            seo_title="Keep SEO title",
            meta_description="Keep meta",
            page_h1="Keep H1",
            seo_faqs=[{"question": "Keep question?", "answer": "Keep answer."}],
            seo_index_mode=Product.SEO_INDEX_NOINDEX,
        )
        product.related_products.add(related)
        offer = Offer.objects.create(product=product, seller=self.seller)
        tier = PricingTier.objects.create(
            offer=offer,
            tier_name="Daily price",
            unit_price="100000",
            minimum_quantity=1,
        )
        previous_verified_at = tier.price_verified_at
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "product_id": product.id,
                "seller_id": self.seller.id,
                "price": "110000",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        product.refresh_from_db()
        tier.refresh_from_db()
        self.assertEqual(product.short_description, "Keep short")
        self.assertEqual(product.description, "<p>Keep full content</p>")
        self.assertEqual(product.seo_title, "Keep SEO title")
        self.assertEqual(product.meta_description, "Keep meta")
        self.assertEqual(product.page_h1, "Keep H1")
        self.assertEqual(product.seo_faqs, [{"question": "Keep question?", "answer": "Keep answer."}])
        self.assertEqual(product.seo_index_mode, Product.SEO_INDEX_NOINDEX)
        self.assertEqual(list(product.related_products.values_list("id", flat=True)), [related.id])
        self.assertGreaterEqual(tier.price_verified_at, previous_verified_at)

    def test_pricing_tier_api_reconfirms_an_unchanged_price(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = Product.objects.create(category=category, name="Seller price verification product")
        offer = Offer.objects.create(product=product, seller=self.seller)
        tier = PricingTier.objects.create(
            offer=offer,
            tier_name="Daily price",
            unit_price="100000",
            minimum_quantity=1,
        )
        old_verified_at = timezone.now() - datetime.timedelta(days=2)
        PricingTier.objects.filter(pk=tier.pk).update(price_verified_at=old_verified_at)
        self.client.force_authenticate(self.admin)

        res = self.client.patch(
            f"/api/pricing-tiers/{tier.id}/",
            {"unit_price": "100000"},
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        tier.refresh_from_db()
        self.assertGreater(tier.price_verified_at, old_verified_at)

    def test_product_summary_exposes_effective_index_state(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        low_information = Product.objects.create(category=category, name="Low")
        forced_index = Product.objects.create(
            category=category,
            name="Forced index product",
            seo_index_mode=Product.SEO_INDEX_INDEX,
        )
        inactive = Product.objects.create(
            category=category,
            name="Inactive forced index product",
            seo_index_mode=Product.SEO_INDEX_INDEX,
            is_active=False,
        )

        res = self.client.get("/api/products-summary/", {"page_size": 1000})

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        index_states = {item["id"]: item["is_indexable"] for item in res.data["results"]}
        self.assertFalse(index_states[low_information.id])
        self.assertTrue(index_states[forced_index.id])
        self.assertFalse(index_states[inactive.id])

    def test_admin_can_manage_category_visual_and_navigation_visibility(self):
        media_root = tempfile.mkdtemp(prefix="category-icons-")
        self.addCleanup(shutil.rmtree, media_root, True)
        category = ProductCategory.objects.get(code="sheet")
        self.client.force_authenticate(self.admin)

        with self.settings(MEDIA_ROOT=media_root):
            visual = self.client.patch(
                f"/api/categories/{category.id}/",
                {
                    "icon_key": "sheet",
                    "icon_image": make_test_category_icon(),
                    "show_in_navigation": False,
                },
                format="multipart",
            )

            self.assertEqual(visual.status_code, status.HTTP_200_OK, visual.data)
            self.assertEqual(visual.data["icon_key"], "sheet")
            self.assertFalse(visual.data["show_in_navigation"])
            self.assertIn("products/category-icons/", visual.data["icon_image"])
            category.refresh_from_db()
            uploaded_name = category.icon_image.name
            self.assertTrue(category.icon_image.storage.exists(uploaded_name))

            invalid = self.client.patch(
                f"/api/categories/{category.id}/",
                {"icon_image": make_test_category_icon("category.gif", "GIF")},
                format="multipart",
            )
            self.assertEqual(invalid.status_code, status.HTTP_400_BAD_REQUEST, invalid.data)
            category.refresh_from_db()
            self.assertEqual(category.icon_image.name, uploaded_name)

            removed = self.client.patch(
                f"/api/categories/{category.id}/",
                {"icon_image": None},
                format="json",
            )
            self.assertEqual(removed.status_code, status.HTTP_200_OK, removed.data)
            self.assertIsNone(removed.data["icon_image"])
            self.assertFalse(category.icon_image.storage.exists(uploaded_name))

    def test_admin_product_price_defaults_to_kilogram_when_basis_is_omitted(self):
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "ورق سیاه بدون مبنای قیمت",
                "category_code": "sheet",
                "seller_id": self.seller.id,
                "price": "140000",
                "steel_grade": "ST37",
                "manufacturing_process": "coil",
                "surface_finish": "black",
                "factory": "mobarakeh",
                "thickness_mm": "2",
                "width_mm": "1250",
                "province": "اصفهان",
                "city": "اصفهان",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        tier = Product.objects.get(id=res.data["product_id"]).offers.get(seller=self.seller).pricing_tiers.get()
        self.assertEqual(tier.price_basis, "kg")
        self.assertEqual(str(tier.unit_price), "140000.00")

    def test_admin_product_ignores_unit_only_pricing_dimensions(self):
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "رول سیاه ساده بدون شرط ابعاد",
                "category_code": "sheet",
                "seller_id": self.seller.id,
                "price": "98000",
                "price_basis": "kg",
                "dimension_width_mm": "mm",
                "dimension_length_mm": "برای رول خالی بماند",
                "steel_grade": "ST37",
                "manufacturing_process": "coil",
                "surface_finish": "black",
                "factory": "mobarakeh",
                "sales_mode": "coil_full",
                "thickness_mm": "15",
                "width_mm": "1500",
                "province": "اصفهان",
                "city": "اصفهان",
                "address": "انبار",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        tier = Product.objects.get(id=res.data["product_id"]).offers.get(seller=self.seller).pricing_tiers.get()
        self.assertIsNone(tier.dimension_width_mm)
        self.assertIsNone(tier.dimension_length_mm)

    def test_admin_product_accepts_blank_pricing_dimensions(self):
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "رول سیاه بدون ابعاد شرط قیمت",
                "category_code": "sheet",
                "seller_id": self.seller.id,
                "price": "98000",
                "price_basis": "kg",
                "dimension_width_mm": "",
                "dimension_length_mm": "",
                "steel_grade": "ST37",
                "manufacturing_process": "coil",
                "surface_finish": "black",
                "factory": "mobarakeh",
                "sales_mode": "coil_full",
                "thickness_mm": "15",
                "width_mm": "1500",
                "length_mm": "",
                "diameter_mm": "",
                "province": "اصفهان",
                "city": "اصفهان",
                "address": "انبار",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        product = Product.objects.get(id=res.data["product_id"])
        self.assertIsNone(product.specifications.length_mm)
        self.assertIsNone(product.specifications.diameter_mm)
        tier = product.offers.get(seller=self.seller).pricing_tiers.get()
        self.assertIsNone(tier.dimension_width_mm)
        self.assertIsNone(tier.dimension_length_mm)

    def test_admin_product_returns_persian_error_for_invalid_numeric_value(self):
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "رول سیاه با شرط عددی نامعتبر",
                "category_code": "sheet",
                "seller_id": self.seller.id,
                "price": "98000",
                "price_basis": "kg",
                "dimension_width_mm": "#4",
                "dimension_length_mm": "",
                "steel_grade": "ST37",
                "manufacturing_process": "coil",
                "surface_finish": "black",
                "factory": "mobarakeh",
                "sales_mode": "coil_full",
                "thickness_mm": "15",
                "width_mm": "1500",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("عرض شرط قیمت باید عدد باشد", str(res.data["detail"]))

    def test_admin_can_update_origin_without_price_change(self):
        category = ProductCategory.objects.get(code="sheet-acid-washed")
        product = Product.objects.create(
            category=category,
            name="ورق اسید مبارکه",
            short_description="ورق اسید مبارکه",
            description="ورق اسید مبارکه",
            is_active=True,
        )
        offer = Offer.objects.create(product=product, seller=self.seller, is_active=True)
        PricingTier.objects.create(
            offer=offer,
            tier_name="قیمت روز",
            unit_price="45000",
            minimum_quantity=1,
        )
        DeliveryLocation.objects.create(
            offer=offer,
            province="اصفهان",
            city="مبارکه",
            address="کارخانه",
        )
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "product_id": product.id,
                "name": product.name,
                "seller_id": self.seller.id,
                "province": "تهران",
                "city": "تهران",
                "address": "انبار",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        self.assertFalse(res.data["updated_price"])
        offer.refresh_from_db()
        delivery = offer.delivery_options.get()
        self.assertEqual(delivery.province, "تهران")
        self.assertEqual(delivery.city, "تهران")
        self.assertEqual(delivery.address, "انبار")
        self.assertEqual(str(offer.pricing_tiers.get(tier_name="قیمت روز").unit_price), "45000.00")

    def test_admin_blank_price_clears_existing_price_and_marks_inquiry(self):
        category = ProductCategory.objects.get(code="sheet-acid-washed")
        product = Product.objects.create(
            category=category,
            name="Blank price product",
            short_description="Blank price product",
            description="Blank price product",
            availability_status=Product.AVAILABILITY_IN_STOCK,
            is_active=True,
        )
        offer = Offer.objects.create(product=product, seller=self.seller, is_active=True)
        PricingTier.objects.create(
            offer=offer,
            tier_name="Daily price",
            unit_price="45000",
            minimum_quantity=1,
        )
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "product_id": product.id,
                "name": product.name,
                "seller_id": self.seller.id,
                "price": "",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        self.assertTrue(res.data["updated_price"])
        self.assertTrue(res.data["cleared_price"])
        product.refresh_from_db()
        self.assertEqual(product.availability_status, Product.AVAILABILITY_INQUIRY)
        self.assertFalse(PricingTier.objects.filter(offer__product=product).exists())

    def test_admin_bulk_price_blank_clears_price_only(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = Product.objects.create(
            category=category,
            name="Bulk blank price product",
            short_description="Keep short",
            description="Keep description",
            availability_status=Product.AVAILABILITY_IN_STOCK,
            is_active=True,
        )
        ProductSpecification.objects.create(
            product=product,
            material_type="sheet",
            steel_grade="ST37",
            surface_finish="black",
            manufacturing_process="sheet",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm="2",
            width_mm="1000",
            length_mm="6000",
        )
        offer = Offer.objects.create(product=product, seller=self.seller, is_active=True)
        PricingTier.objects.create(
            offer=offer,
            tier_name="Daily price",
            unit_price="45000",
            minimum_quantity=1,
        )
        DeliveryLocation.objects.create(
            offer=offer,
            incoterm="EXW",
            country="Iran",
            province="Isfahan",
            city="Mobarakeh",
            address="Warehouse",
        )
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-bulk-prices/",
            {"prices": [{"product_id": product.id, "price": ""}]},
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        self.assertEqual(res.data["updated_prices"], 1)
        product.refresh_from_db()
        self.assertEqual(product.name, "Bulk blank price product")
        self.assertEqual(product.description, "Keep description")
        self.assertEqual(product.availability_status, Product.AVAILABILITY_INQUIRY)
        self.assertEqual(product.specifications.steel_grade, "ST37")
        delivery = offer.delivery_options.get()
        self.assertEqual(delivery.city, "Mobarakeh")
        self.assertFalse(PricingTier.objects.filter(offer__product=product).exists())

    def test_admin_can_create_coil_without_length_or_price(self):
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "رول سیاه مبارکه بدون قیمت",
                "category_code": "sheet",
                "steel_grade": "ST37",
                "manufacturing_process": "coil",
                "surface_finish": "black",
                "factory": "mobarakeh",
                "thickness_mm": "2",
                "width_mm": "1250",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        product = Product.objects.get(id=res.data["product_id"])
        self.assertEqual(product.specifications.manufacturing_process, "coil")
        self.assertIsNone(product.specifications.length_mm)
        self.assertFalse(product.offers.exists())

    def test_admin_product_rejects_unknown_factory_option(self):
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "ورق کارخانه نامعتبر",
                "category_code": "sheet",
                "steel_grade": "ST37",
                "manufacturing_process": "sheet",
                "surface_finish": "black",
                "factory": "unknown_factory",
                "thickness_mm": "2",
                "width_mm": "1250",
                "length_mm": "6000",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("factory", res.data["detail"])

    def test_admin_bulk_xlsx_updates_existing_product_price(self):
        from openpyxl import Workbook

        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = Product.objects.create(
            category=category,
            name="ورق سیاه فایل اکسل",
            short_description="ورق سیاه فایل اکسل",
            description="ورق سیاه فایل اکسل",
        )
        ProductSpecification.objects.create(
            product=product,
            material_type="sheet",
            steel_grade="ST37",
            surface_finish="black",
            manufacturing_process="sheet",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm="2",
            width_mm="1000",
            length_mm="6000",
        )
        offer = Offer.objects.create(product=product, seller=self.seller)
        PricingTier.objects.create(
            offer=offer,
            tier_name="قیمت روز",
            unit_price="41000",
            minimum_quantity=1,
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["شناسه محصول", "شناسه فروشنده", "قیمت جدید", "استان", "شهر"])
        sheet.append([product.id, self.seller.id, 47000, "اصفهان", "مبارکه"])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        upload = SimpleUploadedFile(
            "prices.xlsx",
            payload.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        self.client.force_authenticate(self.admin)
        res = self.client.post(
            "/api/products/admin-bulk-upsert/",
            {"file": upload},
            format="multipart",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        self.assertEqual(res.data["updated_prices"], 1)
        tier = PricingTier.objects.get(offer=offer, tier_name="قیمت روز")
        tier.refresh_from_db()
        self.assertEqual(str(tier.unit_price), "47000.00")
        self.assertEqual(offer.delivery_options.get().province, "اصفهان")
        # رگرسیون: آپلودِ قیمت (بدونِ ستونِ نام) نباید نامِ محصول را خراب کند.
        product.refresh_from_db()
        self.assertEqual(product.name, "ورق سیاه فایل اکسل")
        self.assertNotEqual(product.name, "محصول فولادی")

    def test_admin_bulk_xlsx_price_update_preserves_product_page_content(self):
        from openpyxl import Workbook

        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        related = Product.objects.create(category=category, name="Excel related product")
        product = Product.objects.create(
            category=category,
            name="Excel SEO preservation product",
            short_description="Preserve Excel short description",
            description="<p>Preserve Excel full product content.</p>",
            seo_title="Preserve Excel SEO title",
            meta_description="Preserve Excel meta description",
            page_h1="Preserve Excel H1",
            seo_faqs=[{"question": "Preserve question?", "answer": "Preserve answer."}],
            seo_index_mode=Product.SEO_INDEX_NOINDEX,
        )
        product.related_products.add(related)
        offer = Offer.objects.create(product=product, seller=self.seller)
        tier = PricingTier.objects.create(
            offer=offer,
            tier_name="Daily price",
            unit_price="41000",
            minimum_quantity=1,
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["product_id", "seller_id", "price", "tier_name"])
        sheet.append([product.id, self.seller.id, 47000, "Daily price"])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        upload = SimpleUploadedFile(
            "prices.xlsx",
            payload.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        self.client.force_authenticate(self.admin)
        res = self.client.post(
            "/api/products/admin-bulk-upsert/",
            {"file": upload},
            format="multipart",
        )

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        product.refresh_from_db()
        tier.refresh_from_db()
        self.assertEqual(str(tier.unit_price), "47000.00")
        self.assertIsNotNone(tier.price_verified_at)
        self.assertEqual(product.short_description, "Preserve Excel short description")
        self.assertEqual(product.description, "<p>Preserve Excel full product content.</p>")
        self.assertEqual(product.seo_title, "Preserve Excel SEO title")
        self.assertEqual(product.meta_description, "Preserve Excel meta description")
        self.assertEqual(product.page_h1, "Preserve Excel H1")
        self.assertEqual(product.seo_faqs, [{"question": "Preserve question?", "answer": "Preserve answer."}])
        self.assertEqual(product.seo_index_mode, Product.SEO_INDEX_NOINDEX)
        self.assertEqual(list(product.related_products.values_list("id", flat=True)), [related.id])

    def test_admin_bulk_price_update_preserves_name_without_name_column(self):
        from openpyxl import Workbook

        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = Product.objects.create(
            category=category,
            name="ورق ST37 ویژهٔ من",
            short_description="x",
            description="x",
        )
        offer = Offer.objects.create(product=product, seller=self.seller)
        PricingTier.objects.create(offer=offer, tier_name="قیمت روز", unit_price="100000", minimum_quantity=1)

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["شناسه محصول", "شناسه فروشنده", "قیمت جدید"])
        sheet.append([product.id, self.seller.id, 123000])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        upload = SimpleUploadedFile(
            "prices.xlsx", payload.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        self.client.force_authenticate(self.admin)
        res = self.client.post("/api/products/admin-bulk-upsert/", {"file": upload}, format="multipart")

        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        product.refresh_from_db()
        self.assertEqual(product.name, "ورق ST37 ویژهٔ من")  # دست‌نخورده

    def test_admin_can_download_product_import_template(self):
        self.client.force_authenticate(self.admin)

        res = self.client.get("/api/products/admin-import-template/")

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("kavehmetal-products-template.xlsx", res["Content-Disposition"])

    def test_admin_products_list_includes_inactive_products_with_filters(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        active = Product.objects.create(
            category=category,
            name="محصول فعال ادمین",
            short_description="محصول فعال ادمین",
            description="محصول فعال ادمین",
            is_active=True,
        )
        inactive = Product.objects.create(
            category=category,
            name="محصول غیرفعال ادمین",
            short_description="محصول غیرفعال ادمین",
            description="محصول غیرفعال ادمین",
            is_active=False,
        )
        self.client.force_authenticate(self.admin)

        default_res = self.client.get("/api/products/admin-products/")
        all_res = self.client.get("/api/products/admin-products/", {"is_active": "all"})
        inactive_res = self.client.get("/api/products/admin-products/", {"is_active": "false"})

        self.assertEqual(default_res.status_code, status.HTTP_200_OK)
        default_ids = {item["id"] for item in default_res.data["results"]}
        self.assertIn(active.id, default_ids)
        self.assertIn(inactive.id, default_ids)
        self.assertEqual(all_res.status_code, status.HTTP_200_OK)
        all_ids = {item["id"] for item in all_res.data["results"]}
        self.assertIn(active.id, all_ids)
        self.assertIn(inactive.id, all_ids)
        self.assertEqual(inactive_res.status_code, status.HTTP_200_OK)
        inactive_ids = {item["id"] for item in inactive_res.data["results"]}
        self.assertIn(inactive.id, inactive_ids)
        self.assertNotIn(active.id, inactive_ids)

    def test_admin_can_deactivate_reactivate_and_activity_is_logged(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = Product.objects.create(
            category=category,
            name="محصول وضعیت ادمین",
            short_description="محصول وضعیت ادمین",
            description="محصول وضعیت ادمین",
            is_active=True,
        )
        self.client.force_authenticate(self.admin)

        deactivated = self.client.post(f"/api/products/{product.id}/deactivate/")
        product.refresh_from_db()
        activity = self.client.get("/api/v1/admin/dashboard/activities/", {"limit": 20})
        is_deactivated = not product.is_active
        activated = self.client.post(f"/api/products/{product.id}/activate/")
        product.refresh_from_db()

        self.assertEqual(deactivated.status_code, status.HTTP_200_OK, deactivated.data)
        self.assertTrue(is_deactivated)
        self.assertTrue(ProductAuditLog.objects.filter(product=product, action="PRODUCT_DEACTIVATED").exists())
        self.assertEqual(activity.status_code, status.HTTP_200_OK)
        self.assertTrue(
            any(item["kind"] == "PRODUCT_AUDIT" for item in activity.data["results"])
        )
        self.assertEqual(activated.status_code, status.HTTP_200_OK, activated.data)
        self.assertTrue(product.is_active)
        self.assertTrue(ProductAuditLog.objects.filter(product=product, action="PRODUCT_ACTIVATED").exists())

    def test_admin_can_delete_product_and_keep_audit_log(self):
        category = ProductCategory.objects.get(code="sheet-black-mobarakeh")
        product = Product.objects.create(
            category=category,
            name="محصول حذف ادمین",
            short_description="محصول حذف ادمین",
            description="محصول حذف ادمین",
            is_active=True,
        )
        product_id = product.id
        self.client.force_authenticate(self.admin)

        res = self.client.delete(f"/api/products/{product_id}/")

        self.assertEqual(res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Product.objects.filter(id=product_id).exists())
        log = ProductAuditLog.objects.get(action="PRODUCT_DELETED", product_name="محصول حذف ادمین")
        self.assertIsNone(log.product_id)
        self.assertEqual(log.payload["product_id"], product_id)

    def test_admin_taxonomy_create_actions_are_logged(self):
        self.client.force_authenticate(self.admin)
        sheet = ProductCategory.objects.get(code="sheet")
        black = ProductAttributeOption.objects.get(group="surface_finish", value="black")

        category_res = self.client.post(
            "/api/categories/",
            {
                "name": "ورق تست لاگ",
                "code": "sheet-log-test",
                "parent": sheet.id,
                "product_kind": "sheet",
                "spec_defaults": {"material_type": "sheet", "surface_finish": "black"},
                "required_spec_fields": ["steel_grade", "thickness_mm", "width_mm"],
                "sort_order": 999,
                "is_active": True,
            },
            format="json",
        )
        option_res = self.client.post(
            "/api/attribute-options/",
            {
                "group": "steel_grade",
                "value": "LOGTEST",
                "label": "گرید تست لاگ",
                "product_kind": "sheet",
                "parent": black.id,
                "sort_order": 999,
                "is_active": True,
            },
            format="json",
        )

        self.assertEqual(category_res.status_code, status.HTTP_201_CREATED, category_res.data)
        self.assertEqual(option_res.status_code, status.HTTP_201_CREATED, option_res.data)
        self.assertTrue(
            ProductAuditLog.objects.filter(
                action="PRODUCT_CATEGORY_CREATED",
                product_name="ورق تست لاگ",
            ).exists()
        )
        self.assertTrue(
            ProductAuditLog.objects.filter(
                action="TAXONOMY_OPTION_CREATED",
                product_name="گرید تست لاگ",
            ).exists()
        )

    def test_admin_can_create_and_use_custom_product_kind(self):
        self.client.force_authenticate(self.admin)
        category_res = self.client.post(
            "/api/categories/",
            {
                "name": "نبشی سفارشی",
                "code": "angle-custom",
                "product_kind": "angle",
                "spec_defaults": {"material_type": "angle"},
                "required_spec_fields": [],
                "sort_order": 900,
                "is_active": True,
            },
            format="json",
        )
        option_res = self.client.post(
            "/api/attribute-options/",
            {
                "group": "steel_grade",
                "value": "ST37-ANGLE",
                "label": "ST37 نبشی",
                "product_kind": "angle",
                "sort_order": 1,
                "is_active": True,
            },
            format="json",
        )
        mismatched_child = self.client.post(
            "/api/categories/",
            {
                "name": "زیرگروه ناسازگار نبشی",
                "code": "angle-invalid-child",
                "parent": category_res.data.get("id"),
                "product_kind": "sheet",
                "spec_defaults": {"material_type": "sheet"},
                "is_active": True,
            },
            format="json",
        )
        invalid_kind = self.client.post(
            "/api/categories/",
            {
                "name": "نوع با کد نامعتبر",
                "code": "invalid-kind-code",
                "product_kind": "نبشی",
                "is_active": True,
            },
            format="json",
        )
        product_res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "نبشی ۵۰ در ۵۰ تست",
                "category_code": "angle-custom",
                "seller_id": self.seller.id,
                "price": "52000",
                "steel_grade": "ST37-ANGLE",
                "thickness_mm": "5",
                "width_mm": "50",
                "height_mm": "50",
                "length_mm": "6000",
            },
            format="json",
        )

        self.assertEqual(category_res.status_code, status.HTTP_201_CREATED, category_res.data)
        self.assertEqual(option_res.status_code, status.HTTP_201_CREATED, option_res.data)
        self.assertEqual(mismatched_child.status_code, status.HTTP_400_BAD_REQUEST, mismatched_child.data)
        self.assertIn("product_kind", mismatched_child.data)
        self.assertEqual(invalid_kind.status_code, status.HTTP_400_BAD_REQUEST, invalid_kind.data)
        self.assertIn("product_kind", invalid_kind.data)
        self.assertEqual(product_res.status_code, status.HTTP_200_OK, product_res.data)
        product = Product.objects.get(pk=product_res.data["product_id"])
        self.assertEqual(product.category.product_kind, "angle")
        self.assertEqual(product.specifications.material_type, "angle")
        self.assertEqual(product.specifications.steel_grade, "ST37-ANGLE")
        self.assertEqual(str(product.specifications.height_mm), "50.00")

        public_categories = self.client.get("/api/categories/active-with-products/")
        codes = {item["code"] for item in public_categories.data["results"]}
        self.assertIn("angle-custom", codes)

    def test_admin_can_deactivate_reactivate_category_and_public_list_hides_it(self):
        category = ProductCategory.objects.get(code="sheet-black")
        Product.objects.create(
            category=category,
            name="ورق دسته غیرفعال",
            short_description="ورق دسته غیرفعال",
            description="ورق دسته غیرفعال",
            is_active=True,
        )
        before = self.client.get("/api/categories/active-with-products/")
        self.assertIn(category.id, {item["id"] for item in before.data["results"]})
        self.client.force_authenticate(self.admin)

        deactivated = self.client.post(f"/api/categories/{category.id}/deactivate/")
        public_after = self.client.get("/api/categories/active-with-products/")
        inactive_admin = self.client.get("/api/categories/", {"is_active": "false"})
        activated = self.client.post(f"/api/categories/{category.id}/activate/")

        self.assertEqual(deactivated.status_code, status.HTTP_200_OK, deactivated.data)
        self.assertFalse(deactivated.data["is_active"])
        self.assertNotIn(category.id, {item["id"] for item in public_after.data["results"]})
        self.assertIn(category.id, {item["id"] for item in inactive_admin.data["results"]})
        self.assertEqual(activated.status_code, status.HTTP_200_OK, activated.data)
        self.assertTrue(activated.data["is_active"])
        self.assertTrue(
            ProductAuditLog.objects.filter(
                action="PRODUCT_CATEGORY_DEACTIVATED",
                product_name=category.name,
            ).exists()
        )
        self.assertTrue(
            ProductAuditLog.objects.filter(
                action="PRODUCT_CATEGORY_ACTIVATED",
                product_name=category.name,
            ).exists()
        )

    def test_admin_can_deactivate_option_and_public_lists_hide_it(self):
        option = ProductAttributeOption.objects.get(group="surface_finish", value="black")
        self.client.force_authenticate(self.admin)

        deactivated = self.client.post(f"/api/attribute-options/{option.id}/deactivate/")
        inactive_admin = self.client.get(
            "/api/attribute-options/",
            {"group": "surface_finish", "is_active": "false"},
        )
        self.client.force_authenticate(None)
        inactive_public = self.client.get(
            "/api/attribute-options/",
            {"group": "surface_finish", "is_active": "false"},
        )

        self.assertEqual(deactivated.status_code, status.HTTP_200_OK, deactivated.data)
        self.assertFalse(deactivated.data["is_active"])
        self.assertIn(option.id, {item["id"] for item in inactive_admin.data["results"]})
        self.assertNotIn(option.id, {item["id"] for item in inactive_public.data["results"]})
        self.assertTrue(
            ProductAuditLog.objects.filter(
                action="TAXONOMY_OPTION_DEACTIVATED",
                product_name=option.label,
            ).exists()
        )

    def test_inactive_option_is_rejected_for_new_product_writes(self):
        ProductAttributeOption.objects.filter(group="surface_finish", value="black").update(is_active=False)
        self.client.force_authenticate(self.admin)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "ورق با گزینه غیرفعال",
                "category_code": "sheet",
                "seller_id": self.seller.id,
                "price": "45000",
                "steel_grade": "ST37",
                "manufacturing_process": "sheet",
                "surface_finish": "black",
                "factory": "mobarakeh",
                "cut_type": "cut",
                "thickness_mm": "2",
                "width_mm": "1000",
                "length_mm": "6000",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("surface_finish", str(res.data))

    def test_taxonomy_usage_counts_are_returned(self):
        category = ProductCategory.objects.get(code="sheet-black")
        option = ProductAttributeOption.objects.get(group="surface_finish", value="black")
        product = Product.objects.create(
            category=category,
            name="ورق شمارش",
            short_description="ورق شمارش",
            description="ورق شمارش",
            is_active=True,
        )
        ProductSpecification.objects.create(
            product=product,
            manufacturing_process="sheet",
            surface_finish="black",
            steel_grade="ST37",
            factory="mobarakeh",
            cut_type="cut",
            thickness_mm=2,
            width_mm=1000,
            length_mm=6000,
        )
        self.client.force_authenticate(self.admin)

        category_res = self.client.get(f"/api/categories/{category.id}/")
        option_res = self.client.get(f"/api/attribute-options/{option.id}/")

        self.assertEqual(category_res.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(category_res.data["product_count"], 1)
        self.assertGreaterEqual(category_res.data["active_product_count"], 1)
        self.assertEqual(option_res.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(option_res.data["product_count"], 1)
        self.assertGreaterEqual(option_res.data["active_product_count"], 1)

    def test_non_admin_cannot_use_admin_product_import(self):
        user = User.objects.create_user(username="regular_user", password="pass1234")
        self.client.force_authenticate(user)

        res = self.client.post(
            "/api/products/admin-upsert/",
            {
                "name": "نباید ثبت شود",
                "category_code": "sheet-black-mobarakeh",
                "seller_id": self.seller.id,
                "price": "45000",
            },
            format="json",
        )

        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)
