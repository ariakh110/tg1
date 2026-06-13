import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .models import (
    DeliveryLocation,
    Offer,
    PricingBasis,
    PricingTier,
    Product,
    ProductAttributeOption,
    ProductCategory,
    ProductSpecification,
    Seller,
)
from .serializers import ProductSpecificationSerializer


HEADER_ALIASES = {
    "id": "product_id",
    "productid": "product_id",
    "product_id": "product_id",
    "شناسه": "product_id",
    "شناسه محصول": "product_id",
    "شناسه کالا": "product_id",
    "slug": "slug",
    "اسلاگ": "slug",
    "name": "name",
    "title": "name",
    "product": "name",
    "product_name": "name",
    "نام": "name",
    "نام کالا": "name",
    "عنوان کالا": "name",
    "کالا": "name",
    "short_description": "short_description",
    "توضیح کوتاه": "short_description",
    "description": "description",
    "توضیحات": "description",
    "category_id": "category_id",
    "شناسه دسته": "category_id",
    "category_code": "category_code",
    "code": "category_code",
    "کد دسته": "category_code",
    "کد دسته بندی": "category_code",
    "category": "category_name",
    "category_name": "category_name",
    "دسته": "category_name",
    "دسته بندی": "category_name",
    "نام دسته": "category_name",
    "seller_id": "seller_id",
    "شناسه فروشنده": "seller_id",
    "seller": "seller_name",
    "seller_name": "seller_name",
    "seller_company": "seller_name",
    "فروشنده": "seller_name",
    "نام فروشنده": "seller_name",
    "price": "price",
    "unit_price": "price",
    "price_basis": "price_basis",
    "pricing_basis": "price_basis",
    "unit": "price_basis",
    "basis": "price_basis",
    "مبنای قیمت": "price_basis",
    "واحد قیمت": "price_basis",
    "condition_label": "condition_label",
    "pricing_condition": "condition_label",
    "dimension_label": "condition_label",
    "شرط قیمت": "condition_label",
    "عنوان ابعاد": "condition_label",
    "dimension_width": "dimension_width_mm",
    "dimension_width_mm": "dimension_width_mm",
    "pricing_width": "dimension_width_mm",
    "عرض قیمت": "dimension_width_mm",
    "dimension_length": "dimension_length_mm",
    "dimension_length_mm": "dimension_length_mm",
    "pricing_length": "dimension_length_mm",
    "طول قیمت": "dimension_length_mm",
    "قیمت": "price",
    "قیمت روز": "price",
    "قیمت جدید": "price",
    "قیمت تومان": "price",
    "tier_name": "tier_name",
    "نام قیمت": "tier_name",
    "material_type": "material_type",
    "نوع محصول": "material_type",
    "steel_grade": "steel_grade",
    "grade": "steel_grade",
    "alloy": "steel_grade",
    "گرید": "steel_grade",
    "آلیاژ": "steel_grade",
    "thickness": "thickness_mm",
    "thickness_mm": "thickness_mm",
    "ضخامت": "thickness_mm",
    "width": "width_mm",
    "width_mm": "width_mm",
    "عرض": "width_mm",
    "length": "length_mm",
    "length_mm": "length_mm",
    "طول": "length_mm",
    "height": "height_mm",
    "height_mm": "height_mm",
    "ارتفاع": "height_mm",
    "diameter": "diameter_mm",
    "diameter_mm": "diameter_mm",
    "قطر": "diameter_mm",
    "weight": "weight_kg_per_unit",
    "weight_kg_per_unit": "weight_kg_per_unit",
    "وزن": "weight_kg_per_unit",
    "surface_finish": "surface_finish",
    "نوع سطح": "surface_finish",
    "manufacturing_process": "manufacturing_process",
    "process": "manufacturing_process",
    "فرایند تولید": "manufacturing_process",
    "نوع تولید": "manufacturing_process",
    "نوع ورق/رول": "manufacturing_process",
    "رول یا ورق": "manufacturing_process",
    "factory": "factory",
    "manufacturer": "factory",
    "mill": "factory",
    "کارخانه": "factory",
    "تولید کننده": "factory",
    "cut_type": "cut_type",
    "fabrication": "cut_type",
    "نوع برش": "cut_type",
    "فابریک یا برش خورده": "cut_type",
    "availability_status": "availability_status",
    "purchase_terms": "purchase_terms",
    "terms": "purchase_terms",
    "sales_mode": "sales_mode",
    "head_tail_policy": "head_tail_policy",
    "وضعیت موجودی": "availability_status",
    "موجودی": "availability_status",
    "country": "country",
    "کشور": "country",
    "province": "province",
    "استان": "province",
    "city": "city",
    "شهر": "city",
    "address": "address",
    "warehouse": "address",
    "آدرس": "address",
    "انبار": "address",
    "incoterm": "incoterm",
}

SPEC_FIELDS = {
    "material_type",
    "steel_grade",
    "surface_finish",
    "manufacturing_process",
    "factory",
    "cut_type",
    "thickness_mm",
    "width_mm",
    "length_mm",
    "height_mm",
    "diameter_mm",
    "weight_kg_per_unit",
    "sales_mode",
    "head_tail_policy",
}

NUMERIC_FIELDS = {
    "price",
    "thickness_mm",
    "width_mm",
    "length_mm",
    "height_mm",
    "diameter_mm",
    "weight_kg_per_unit",
    "dimension_width_mm",
    "dimension_length_mm",
}

NUMERIC_FIELD_LABELS = {
    "price": "قیمت واحد روز",
    "thickness_mm": "ضخامت",
    "width_mm": "عرض",
    "length_mm": "طول",
    "height_mm": "ارتفاع",
    "diameter_mm": "قطر",
    "weight_kg_per_unit": "وزن هر واحد",
    "dimension_width_mm": "عرض شرط قیمت",
    "dimension_length_mm": "طول شرط قیمت",
}

PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def option_label(group, value):
    if value in ("", None):
        return ""
    option = ProductAttributeOption.objects.filter(group=group, value=value, is_active=True).first()
    return option.label if option else str(value)


def format_dimension(value):
    if value in ("", None):
        return ""
    number = Decimal(value)
    text = f"{number.normalize():f}"
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def build_product_name(row, category):
    material = option_label("surface_finish", row.get("surface_finish"))
    process = option_label("manufacturing_process", row.get("manufacturing_process"))
    factory = option_label("factory", row.get("factory"))
    thickness = format_dimension(row.get("thickness_mm"))
    width = format_dimension(row.get("width_mm"))
    length = format_dimension(row.get("length_mm"))
    parts = []
    if category:
        parts.append(category.name)
    if material and material not in parts:
        parts.append(material)
    if thickness:
        parts.append(f"{thickness} میل")
    if width and length:
        parts.append(f"{width}x{length}")
    elif width:
        parts.append(f"عرض {width}")
    if process:
        parts.append(process)
    if factory:
        parts.append(factory)
    return " ".join(parts).strip() or "محصول فولادی"


def normalize_header(value):
    text = str(value or "").strip().translate(PERSIAN_DIGITS)
    text = text.replace("\u200c", " ").replace("-", " ").replace("_", " ")
    text = " ".join(text.split()).lower()
    compact = text.replace(" ", "_")
    return HEADER_ALIASES.get(text) or HEADER_ALIASES.get(compact) or compact


def clean_value(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def parse_terms(value):
    if value in ("", None):
        return []
    text = str(value).replace("\r", "\n")
    parts = []
    for chunk in text.replace("؛", "|").replace(";", "|").split("|"):
        for line in chunk.split("\n"):
            item = line.strip()
            if item:
                parts.append(item[:500])
    return parts


def normalize_row(row):
    normalized = {}
    for key, value in row.items():
        canonical = normalize_header(key)
        if canonical:
            normalized[canonical] = clean_value(value)
    return normalized


def parse_decimal(value, field):
    if value in ("", None):
        return None
    text = str(value).strip().translate(PERSIAN_DIGITS)
    if text.lower() == "mm" or "خالی" in text:
        return None
    text = (
        text.replace(",", "")
        .replace("٬", "")
        .replace("،", "")
        .replace("تومان", "")
        .replace("ریال", "")
        .replace("mm", "")
        .replace("MM", "")
        .replace("میلی‌متر", "")
        .replace("میلیمتر", "")
        .strip()
    )
    if text == "":
        return None
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        label = NUMERIC_FIELD_LABELS.get(field, field)
        raise ValueError(f"{label} باید عدد باشد.") from exc


def normalize_price_basis(value):
    value = str(value or PricingBasis.KG).strip().lower()
    aliases = {
        "t": PricingBasis.TON,
        "ton": PricingBasis.TON,
        "tons": PricingBasis.TON,
        "tonne": PricingBasis.TON,
        "تن": PricingBasis.TON,
        "هر تن": PricingBasis.TON,
        "kg": PricingBasis.KG,
        "kilo": PricingBasis.KG,
        "kilogram": PricingBasis.KG,
        "کیلو": PricingBasis.KG,
        "کیلوگرم": PricingBasis.KG,
        "هر کیلو": PricingBasis.KG,
        "sheet": PricingBasis.SHEET,
        "sheets": PricingBasis.SHEET,
        "sheet_count": PricingBasis.SHEET,
        "ورق": PricingBasis.SHEET,
        "هر ورق": PricingBasis.SHEET,
        "تعداد ورق": PricingBasis.SHEET,
    }
    allowed = {PricingBasis.TON, PricingBasis.KG, PricingBasis.SHEET}
    result = aliases.get(value, value)
    return result if result in allowed else PricingBasis.KG


def read_price_file(uploaded_file):
    filename = (getattr(uploaded_file, "name", "") or "").lower()
    if filename.endswith(".csv") or filename.endswith(".txt"):
        raw = uploaded_file.read()
        text = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return [normalize_row(row) for row in reader if any(clean_value(v) for v in row.values())]

    if filename.endswith(".xlsx") or filename.endswith(".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("برای خواندن فایل Excel پکیج openpyxl باید نصب باشد.") from exc

        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = None
        data = []
        for raw_row in rows:
            if not raw_row or not any(clean_value(value) for value in raw_row):
                continue
            if headers is None:
                headers = [normalize_header(value) for value in raw_row]
                continue
            row = {
                headers[index]: raw_row[index]
                for index in range(min(len(headers), len(raw_row)))
                if headers[index]
            }
            normalized = normalize_row(row)
            if any(normalized.values()):
                data.append(normalized)
        return data

    raise ValueError("فقط فایل CSV یا XLSX پشتیبانی می‌شود.")


def resolve_category(row):
    if row.get("category_id"):
        return ProductCategory.objects.get(pk=row["category_id"])
    if row.get("category_code"):
        return ProductCategory.objects.get(code=row["category_code"])
    if row.get("category_name"):
        name = row["category_name"].split("/")[-1].strip()
        return ProductCategory.objects.get(name=name)
    return None


def resolve_seller(row, default_seller_id=None):
    seller_id = row.get("seller_id") or default_seller_id
    if seller_id:
        return Seller.objects.get(pk=seller_id)
    if row.get("seller_name"):
        return Seller.objects.get(company_name=row["seller_name"])
    return None


def find_product(row, category):
    if row.get("product_id"):
        return Product.objects.get(pk=row["product_id"])
    if row.get("slug"):
        product = Product.objects.filter(slug=row["slug"]).first()
        if product:
            return product
    if row.get("name") and category:
        return Product.objects.filter(name=row["name"], category=category).first()
    if row.get("name"):
        return Product.objects.filter(name=row["name"]).first()
    return None


def upsert_product_row(row, *, default_seller_id=None, create_missing=True):
    row = normalize_row(row)
    for field in NUMERIC_FIELDS:
        if field in row:
            row[field] = parse_decimal(row[field], field)

    category = resolve_category(row)
    product = find_product(row, category)
    created_product = False
    generated_name = row.get("name") or build_product_name(row, category)

    if product is None:
        if not create_missing:
            raise ValueError("محصول پیدا نشد و ایجاد محصول جدید غیرفعال است.")
        if category is None:
            raise ValueError("برای ایجاد محصول جدید، دسته‌بندی لازم است.")
        product = Product.objects.create(
            name=generated_name,
            category=category,
            short_description=row.get("short_description") or generated_name,
            description=row.get("description") or row.get("short_description") or generated_name,
            availability_status=row.get("availability_status") or Product.AVAILABILITY_IN_STOCK,
            purchase_terms=parse_terms(row.get("purchase_terms")),
            is_active=True,
        )
        created_product = True
    else:
        changed_fields = []
        if generated_name and product.name != generated_name:
            product.name = generated_name
            changed_fields.append("name")
        if category and product.category_id != category.id:
            product.category = category
            changed_fields.append("category")
        if row.get("short_description"):
            product.short_description = row["short_description"]
            changed_fields.append("short_description")
        if row.get("description"):
            product.description = row["description"]
            changed_fields.append("description")
        if row.get("availability_status") and product.availability_status != row["availability_status"]:
            product.availability_status = row["availability_status"]
            changed_fields.append("availability_status")
        if row.get("purchase_terms") not in ("", None):
            product.purchase_terms = parse_terms(row.get("purchase_terms"))
            changed_fields.append("purchase_terms")
        if changed_fields:
            product.save(update_fields=changed_fields)

    spec_payload = {field: row[field] for field in SPEC_FIELDS if row.get(field) not in ("", None)}
    should_write_spec = bool(spec_payload) or created_product
    if should_write_spec:
        try:
            instance = product.specifications
        except ProductSpecification.DoesNotExist:
            instance = None
        serializer_data = {"product": product.id, **spec_payload}
        serializer = ProductSpecificationSerializer(
            instance,
            data=serializer_data,
            partial=bool(instance),
        )
        serializer.is_valid(raise_exception=True)
        serializer.save()

    price = row.get("price")
    updated_price = False
    has_price = price not in (None, "")
    has_delivery = any(row.get(field) not in ("", None) for field in ("province", "city", "address", "country", "incoterm"))
    offer = None
    if has_price or has_delivery:
        seller = resolve_seller(row, default_seller_id=default_seller_id)
        if seller is None:
            raise ValueError("برای ثبت قیمت یا مبدا بار، فروشنده لازم است.")
        offer, _ = Offer.objects.get_or_create(product=product, seller=seller, defaults={"is_active": True})
    if has_price:
        tier_name = row.get("tier_name") or "قیمت روز"
        price_basis = normalize_price_basis(row.get("price_basis"))
        condition_label = row.get("condition_label") or ""
        dimension_width = row.get("dimension_width_mm")
        dimension_length = row.get("dimension_length_mm")
        tier = PricingTier.objects.filter(offer=offer, tier_name=tier_name).first()
        if tier is None:
            PricingTier.objects.create(
                offer=offer,
                tier_name=tier_name,
                unit_price=price,
                price_basis=price_basis,
                minimum_quantity=1,
                maximum_quantity=None,
                condition_label=condition_label,
                dimension_width_mm=dimension_width,
                dimension_length_mm=dimension_length,
                is_negotiable=False,
            )
        else:
            # فقط قیمت/واحد را همیشه به‌روزرسانی کن؛ شرط و ابعادِ tier را تنها
            # وقتی ردیف مقدار غیرخالی بدهد عوض کن تا آپلودِ «فقط قیمت» داده‌های موجود را پاک نکند.
            tier.unit_price = price
            tier.price_basis = price_basis
            tier.minimum_quantity = tier.minimum_quantity or 1
            update_fields = ["unit_price", "price_basis", "minimum_quantity"]
            if row.get("condition_label") not in ("", None):
                tier.condition_label = condition_label
                update_fields.append("condition_label")
            if row.get("dimension_width_mm") not in ("", None):
                tier.dimension_width_mm = dimension_width
                update_fields.append("dimension_width_mm")
            if row.get("dimension_length_mm") not in ("", None):
                tier.dimension_length_mm = dimension_length
                update_fields.append("dimension_length_mm")
            tier.save(update_fields=update_fields)
        updated_price = True

    if has_delivery and offer is not None:
        delivery = offer.delivery_options.first()
        if delivery is None:
            DeliveryLocation.objects.create(
                offer=offer,
                incoterm=row.get("incoterm") or "EXW",
                country=row.get("country") or "ایران",
                province=row.get("province") or "",
                city=row.get("city") or "",
                address=row.get("address") or "",
            )
        else:
            # روی رکورد موجود فقط فیلدهای ارائه‌شده را بنویس تا خالی‌ها پاک‌کننده نباشند.
            provided = {
                field: row.get(field)
                for field in ("incoterm", "country", "province", "city", "address")
                if row.get(field) not in ("", None)
            }
            if provided:
                for field, value in provided.items():
                    setattr(delivery, field, value)
                delivery.save(update_fields=list(provided.keys()))

    return {
        "product_id": product.id,
        "product_name": product.name,
        "created_product": created_product,
        "updated_price": updated_price,
    }


def bulk_upsert_products(rows, *, default_seller_id=None, create_missing=True):
    created = 0
    updated = 0
    errors = []
    results = []

    for index, row in enumerate(rows, start=2):
        try:
            with transaction.atomic():
                result = upsert_product_row(
                    row,
                    default_seller_id=default_seller_id,
                    create_missing=create_missing,
                )
            created += 1 if result["created_product"] else 0
            updated += 1 if result["updated_price"] else 0
            results.append(result)
        except Exception as exc:
            detail = getattr(exc, "detail", None)
            errors.append({"row": index, "error": detail or str(exc)})

    return {
        "total_rows": len(rows),
        "created_products": created,
        "updated_prices": updated,
        "failed_rows": len(errors),
        "errors": errors,
        "results": results[:20],
    }
